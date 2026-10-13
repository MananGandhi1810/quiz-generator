import os
import subprocess
import time
import zipfile
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

MAX_PAGES = int(os.getenv('MAX_PAGES', '40'))
MAX_TEXT = int(os.getenv('MAX_TEXT_CHARS', '80000'))
SUPPORTED = {'.pdf', '.docx', '.pptx', '.doc', '.ppt'}


class DocumentError(ValueError):
    pass


def run_command(args, timeout=90, deadline=None):
    if deadline is not None:
        timeout = min(timeout, deadline - time.monotonic())
        if timeout <= 0:
            raise DocumentError('Document processing took too long. Try a smaller document.')
    try:
        result = subprocess.run(args, capture_output=True, timeout=timeout, check=False)
    except FileNotFoundError:
        raise DocumentError('Document processing tools are missing. Please run the app with Docker.') from None
    except subprocess.TimeoutExpired:
        raise DocumentError('Document processing took too long. Try a smaller document.') from None
    if result.returncode:
        raise DocumentError('This document could not be processed. It may be corrupt or password-protected.')


def pdf_text(path):
    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(''):
        raise DocumentError('Password-protected PDFs are not supported. Upload an unlocked copy.')
    if len(reader.pages) > MAX_PAGES:
        raise DocumentError(f'Please upload a document with at most {MAX_PAGES} pages.')
    chunks, ocr_pages = [], 0
    deadline = time.monotonic() + 600
    for index, page in enumerate(reader.pages):
        text = page.extract_text() or ''
        if len(''.join(text.split())) < 40:
            prefix = path.parent / f'page-{index}'
            run_command(['pdftoppm', '-f', str(index + 1), '-l', str(index + 1), '-scale-to', '2500', '-singlefile', '-png', str(path), str(prefix)], deadline=deadline)
            output = prefix.with_suffix('.txt')
            run_command(['tesseract', str(prefix.with_suffix('.png')), str(prefix), '-l', os.getenv('OCR_LANGUAGE', 'eng')], deadline=deadline)
            text = output.read_text()
            prefix.with_suffix('.png').unlink(missing_ok=True)
            output.unlink(missing_ok=True)
            ocr_pages += 1
        chunks.append(f'[Page {index + 1}]\n{text}')
    return '\n\n'.join(chunks), ocr_pages


def check_archive(path):
    with zipfile.ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise DocumentError('The expanded document is too large. Try a smaller file.')


def shape_text(shapes):
    parts = []
    for shape in shapes:
        if shape.has_text_frame:
            parts.append(shape.text)
        if shape.has_table:
            parts.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
        if shape.shape_type == 6:  # grouped shapes
            parts.extend(shape_text(shape.shapes))
    return parts


def extract_document(path: Path):
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise DocumentError('Upload a PDF, Word (.doc/.docx), or PowerPoint (.ppt/.pptx) file.')
    try:
        if suffix in {'.doc', '.ppt'}:
            target = 'docx' if suffix == '.doc' else 'pptx'
            run_command(['libreoffice', f'-env:UserInstallation=file://{path.parent}/office-profile', '--headless', '--convert-to', target, '--outdir', str(path.parent), str(path)])
            path = path.with_suffix('.' + target)
            if not path.exists():
                raise DocumentError('Office conversion failed. Try saving the file as DOCX or PPTX.')
            suffix = path.suffix
        ocr_pages = 0
        if suffix == '.pdf':
            text, ocr_pages = pdf_text(path)
        else:
            check_archive(path)
            if suffix == '.docx':
                document = Document(path)
                parts = [paragraph.text for paragraph in document.paragraphs]
                parts.extend(' | '.join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
                text = '\n'.join(parts)
            else:
                presentation = Presentation(path)
                if len(presentation.slides) > MAX_PAGES:
                    raise DocumentError(f'Please upload a presentation with at most {MAX_PAGES} slides.')
                parts = []
                for i, slide in enumerate(presentation.slides):
                    parts.append(f'[Slide {i + 1}]')
                    parts.extend(shape_text(slide.shapes))
                    if slide.has_notes_slide:
                        frame = slide.notes_slide.notes_text_frame
                        if frame is not None:
                            parts.append(frame.text)
                text = '\n'.join(parts)
        # Labels alone should not count as readable notes.
        content = '\n'.join(line for line in text.splitlines() if not line.startswith(('[Page ', '[Slide '))).strip()
        if len(content) < 80:
            raise DocumentError('We could not find enough readable notes. Try a clearer scan or a document with more text.')
        if len(text) > MAX_TEXT:
            raise DocumentError(f'The document contains too much text (limit {MAX_TEXT:,} characters). Split it into smaller sections.')
        return text, ocr_pages
    except DocumentError:
        raise
    except Exception:
        raise DocumentError('This file could not be read. Check that it is a valid, unlocked document.') from None
