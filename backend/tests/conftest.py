import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

NOTES = ('Photosynthesis converts light energy to chemical energy. Chlorophyll absorbs light in chloroplasts. '
         'The light-dependent reactions produce ATP and NADPH. Water is split, releasing oxygen. '
         'The Calvin cycle uses carbon dioxide to produce sugars. Mitochondria perform respiration. '
         'Ribosomes synthesize proteins. The nucleus contains DNA. Osmosis moves water across membranes.')

@pytest.fixture
def documents(tmp_path):
    doc = Document()
    doc.add_paragraph(NOTES)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Cell membrane'
    table.cell(0, 1).text = 'Controls transport'
    doc.save(tmp_path / 'notes.docx')
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4)).text = NOTES
    slide.notes_slide.notes_text_frame.text = 'Speaker note: Cells are the basic unit of life.'
    ppt.save(tmp_path / 'notes.pptx')
    pdf = canvas.Canvas(str(tmp_path / 'notes.pdf'))
    for i, sentence in enumerate(NOTES.split('. ')):
        pdf.drawString(40, 780-i*25, sentence)
    pdf.save()
    return tmp_path
