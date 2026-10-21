import shutil
import pytest
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfWriter
from app.documents import extract_document, DocumentError, run_command
from conftest import NOTES

@pytest.mark.parametrize('extension', ['pdf', 'docx', 'pptx'])
def test_text_documents(documents, extension):
    text, ocr = extract_document(documents / f'notes.{extension}')
    assert 'Photosynthesis' in text
    assert 'Ribosomes' in text
    assert ocr == 0
    if extension == 'pptx': assert 'Speaker note' in text
    if extension == 'docx': assert 'Controls transport' in text


def test_scanned_pdf(tmp_path):
    if not all(shutil.which(tool) for tool in ['tesseract', 'pdftoppm']): pytest.skip('OCR tools unavailable')
    image = Image.new('RGB', (1800, 1000), 'white')
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=30)
    for i, sentence in enumerate(NOTES.split('. ')): draw.text((60, 60+i*80), sentence, fill='black', font=font)
    path = tmp_path / 'scan.pdf'
    image.save(path, 'PDF', resolution=150)
    text, ocr = extract_document(path)
    assert ocr == 1
    assert 'Photosynthesis' in text
    assert 'Ribosomes' in text

@pytest.mark.parametrize('source,target,filter_name', [('docx','doc','MS Word 97'),('pptx','ppt','MS PowerPoint 97')])
def test_legacy_documents(documents, source, target, filter_name):
    if not shutil.which('libreoffice'): pytest.skip('LibreOffice unavailable')
    run_command(['libreoffice', f'-env:UserInstallation=file://{documents}/fixture-profile', '--headless', '--convert-to', f'{target}:{filter_name}', '--outdir', str(documents), str(documents/f'notes.{source}')])
    # Remove modern fixture so extraction must produce it again.
    (documents/f'notes.{source}').unlink()
    text, _ = extract_document(documents/f'notes.{target}')
    assert 'Photosynthesis' in text


def test_locked_pdf(tmp_path):
    writer = PdfWriter(); writer.add_blank_page(width=100, height=100); writer.encrypt('secret')
    path = tmp_path/'locked.pdf'
    writer.write(path)
    with pytest.raises(DocumentError, match='Password-protected'): extract_document(path)


def test_corrupt_pdf(tmp_path):
    path = tmp_path/'broken.pdf'; path.write_bytes(b'not a document')
    with pytest.raises(DocumentError, match='could not be read'): extract_document(path)


def test_text_limit(documents, monkeypatch):
    monkeypatch.setattr('app.documents.MAX_TEXT', 100)
    with pytest.raises(DocumentError, match='too much text'): extract_document(documents/'notes.docx')


@pytest.mark.parametrize('extension', ['pdf', 'pptx'])
def test_documents_over_40_pages(tmp_path, extension):
    path = tmp_path / f'long.{extension}'
    if extension == 'pdf':
        from reportlab.pdfgen import canvas
        pdf = canvas.Canvas(str(path))
        for _ in range(41):
            pdf.drawString(40, 780, NOTES)
            pdf.showPage()
        pdf.save()
    else:
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        for _ in range(41):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4)).text = NOTES
        presentation.save(path)
    text, ocr = extract_document(path)
    label = 'Page' if extension == 'pdf' else 'Slide'
    assert f'[{label} 41]' in text
    assert text.count('Photosynthesis') == 41
    assert ocr == 0


def test_empty_word(tmp_path):
    from docx import Document
    path = tmp_path/'empty.docx'; Document().save(path)
    with pytest.raises(DocumentError, match='enough readable notes'): extract_document(path)
